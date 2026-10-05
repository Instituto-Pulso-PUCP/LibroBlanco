"""Parseo de la matriz CEPLAN/PEDN 2050 (datos/PEDN2050.xlsx).

No es un dominio como projects/publications: no tiene unidades que embeber ni
temas que extraer con un LLM. CEPLAN ya trae su propia taxonomia oficial,
nombrada -- usar un LLM para "redescubrirla" seria peor que usarla tal cual.
Este modulo solo la convierte en un arbol ON -> Tematica -> Sub-tematica con
texto listo para embeber, para comparar contra los temas normalizados de
projects/publications_linked (ver scripts/pipeline_temas/06_ceplan_alignment.py).

Estructura del libro (ver docs/cuantificacion_temas.md para el detalle):
  - Hojas "ON 1".."ON 4 ": la matriz completa (Politica de Estado -> Politica
    Nacional -> EJE -> ON -> Objetivo Especifico -> Tematica/Sub-tematica ->
    Accion Estrategica -> PGG). Aporta texto de enriquecimiento por
    Sub-tematica (Objetivo Especifico + Accion Estrategica).
  - Hoja "Lineas de Inv.": el arbol ON -> Tematica -> Sub-tematica ya
    curado como "oportunidades de colaboracion" -- es la taxonomia PRIMARIA
    para este cruce, construida especificamente para este proposito.
  - Hoja "Lineas Ceplan": el mismo arbol pero solo "trabajos en desarrollo
    por la Direccion" -- taxonomia SECUNDARIA, mas chica, de referencia.

Las dos hojas "Lineas *" traen, ademas de los 4 ON canonicos, dos categorias
transversales ("Estudios relacionados al analisis de futuro",
"...conocimiento integral de la realidad") que no son un ON: se conservan
aparte, no entran en el cruce con ON.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl

REPO_ROOT = Path(__file__).resolve().parents[2]
PEDN_XLSX = REPO_ROOT / "datos" / "PEDN2050.xlsx"

ON_SHEETS = ("ON 1", "ON 2", "ON 3", "ON 4 ")  # ojo: "ON 4 " trae un espacio
LINEAS_SHEETS = {"primary": "Lineas de Inv.", "secondary": "Lineas Ceplan"}
# Los nombres de hoja reales usan tildes; openpyxl es exacto con el nombre.
_SHEET_NAME_MAP = {"Lineas de Inv.": "Líneas de Inv.", "Lineas Ceplan": "Líneas Ceplan"}

_JUNK = {"", "-", "n/a", "na", "null", "none"}


def _clean(v) -> str:
    if v is None:
        return ""
    s = str(v).strip()
    return "" if s.lower() in _JUNK else s


def _is_on(label: str) -> bool:
    return label.upper().replace(" ", "").startswith("ON")


def _load_lineas_sheet(wb, sheet_key: str) -> list[dict]:
    """Lee 'Lineas de Inv.' o 'Lineas Ceplan': 4 columnas (vacia, ON,
    Tematica, Sub-tematica), con forward-fill porque las celdas combinadas
    solo llevan el valor en su primera fila."""
    name = _SHEET_NAME_MAP[LINEAS_SHEETS[sheet_key]]
    ws = wb[name]
    rows = []
    last = [None, None, None]
    # Fila 5 (1-based) es el encabezado "Objetivo Nacional | Tematicas | Sub
    # tematicas"; los datos empiezan en la fila 6.
    for raw in ws.iter_rows(min_row=6, max_col=4, values_only=True):
        cols = [_clean(c) for c in raw[1:4]]
        if not any(cols):
            continue
        for i in range(3):
            if cols[i]:
                last[i] = cols[i]
            cols[i] = last[i]
        on_label, tema, subtema = cols
        if not subtema:
            continue
        rows.append({"on": on_label, "tema": tema, "subtema": subtema})
    return rows


def _load_on_sheet_enrichment(wb) -> dict[tuple[str, str], list[str]]:
    """De las hojas 'ON N', junta el texto de Objetivo Especifico + Accion
    Estrategica de cada (Tematica, Sub-tematica), para enriquecer el nodo
    hoja del arbol de 'Lineas de Inv.' con mas contexto que solo su nombre."""
    enrichment: dict[tuple[str, str], list[str]] = {}
    for sheet_name in ON_SHEETS:
        ws = wb[sheet_name]
        for row in ws.iter_rows(min_row=5, values_only=True):
            if not any(c is not None for c in row):
                continue
            tema = _clean(row[17]) if len(row) > 17 else ""       # Tematicas
            subtema = _clean(row[24]) if len(row) > 24 else ""    # Sub tematicas
            oe = _clean(row[18]) if len(row) > 18 else ""         # Objetivo Especifico
            ae = _clean(row[25]) if len(row) > 25 else ""         # Accion Estrategica
            if not subtema:
                continue
            key = (tema, subtema)
            for text in (oe, ae):
                if text:
                    enrichment.setdefault(key, [])
                    if text not in enrichment[key]:
                        enrichment[key].append(text)
    return enrichment


def build_taxonomy() -> dict:
    """Devuelve {'on': {...}, 'excluded_categories': [...], 'secondary': {...}}.

    ``on``: {on_id: {'label', 'temas': {tema: {'label', 'subtemas': {subtema_id: {...}}}}}}
    Cada sub-tematica trae 'label' y 'text' (label + tematica + enriquecimiento,
    listo para embeber).
    """
    if not PEDN_XLSX.exists():
        raise FileNotFoundError(
            f"No existe {PEDN_XLSX}. Es un insumo de datos/, no se regenera.")
    wb = openpyxl.load_workbook(PEDN_XLSX, read_only=True, data_only=True)

    primary_rows = _load_lineas_sheet(wb, "primary")
    secondary_rows = _load_lineas_sheet(wb, "secondary")
    enrichment = _load_on_sheet_enrichment(wb)

    def to_tree(rows, with_enrichment=False):
        on_tree: dict[str, dict] = {}
        excluded: list[dict] = []
        for r in rows:
            on_label, tema, subtema = r["on"], r["tema"], r["subtema"]
            if not on_label or not _is_on(on_label):
                excluded.append(r)
                continue
            on_id = on_label.split(".")[0].strip().replace(" ", "").upper()
            on_node = on_tree.setdefault(on_id, {"label": on_label, "temas": {}})
            tema_key = tema or "(sin tematica)"
            tema_node = on_node["temas"].setdefault(
                tema_key, {"label": tema_key, "subtemas": {}})
            sub_id = subtema.strip().lower().replace(" ", "-")[:60]
            parts = [subtema]
            if tema_key != "(sin tematica)":
                parts.append(tema_key)
            if with_enrichment:
                parts.extend(enrichment.get((tema, subtema), []))
            tema_node["subtemas"][sub_id] = {
                "label": subtema, "text": ". ".join(parts)}
        return on_tree, excluded

    on_tree, excluded = to_tree(primary_rows, with_enrichment=True)
    secondary_tree, _ = to_tree(secondary_rows, with_enrichment=False)

    return {"on": on_tree, "excluded_categories": excluded, "secondary": secondary_tree}


TRANSLATION_CACHE = REPO_ROOT / "salidas" / "topics" / "ceplan" / "subtematicas_en.json"

_TRANSLATE_SYSTEM = (
    "Traduces al ingles textos de planificacion publica peruana (PEDN 2050, CEPLAN). "
    "Traduccion fiel y natural, sin resumir ni agregar nada. Responde SOLO con un "
    "arreglo JSON de strings, uno por texto de entrada, en el mismo orden.")


def translate_texts(llm, texts: list[str], batch_size: int = 15) -> dict[str, str]:
    """Traduce al ingles los textos de sub-tematica, con cache en disco.

    Por que hace falta: Cohere multilingual v3 da similitudes ~0.15 mas bajas
    entre un texto en ingles y uno en espanol que entre dos del mismo idioma,
    con el mismo contenido. CEPLAN esta en espanol y ~2/3 de las publicaciones
    ligadas en ingles, asi que un piso absoluto de similitud las marcaria como
    "sin relacion" solo por el idioma. Con una version en ingles de cada
    sub-tematica se toma la mejor de las dos similitudes (ver
    scripts/analysis/ceplan_units.py). La cache se indexa por el texto en
    espanol: si CEPLAN cambia una sub-tematica, solo esa se retraduce.
    """
    cache = {}
    if TRANSLATION_CACHE.exists():
        cache = json.loads(TRANSLATION_CACHE.read_text(encoding="utf-8"))
    pending = [t for t in dict.fromkeys(texts) if t not in cache]
    for start in range(0, len(pending), batch_size):
        chunk = pending[start:start + batch_size]
        raw = llm.complete(_TRANSLATE_SYSTEM, json.dumps(chunk, ensure_ascii=False),
                           max_tokens=8000)
        raw = raw.strip()
        raw = raw[raw.find("["):raw.rfind("]") + 1]
        out = json.loads(raw)
        if len(out) != len(chunk):
            raise RuntimeError(f"La traduccion devolvio {len(out)} textos, no {len(chunk)}")
        cache.update(zip(chunk, out))
        TRANSLATION_CACHE.parent.mkdir(parents=True, exist_ok=True)
        TRANSLATION_CACHE.write_text(json.dumps(cache, ensure_ascii=False, indent=1),
                                     encoding="utf-8")
        print(f"  traducidas {min(start + batch_size, len(pending))}/{len(pending)}")
    return {t: cache[t] for t in texts}


if __name__ == "__main__":
    tax = build_taxonomy()
    n_temas = sum(len(o["temas"]) for o in tax["on"].values())
    n_subtemas = sum(len(t["subtemas"]) for o in tax["on"].values() for t in o["temas"].values())
    print(f"ON: {len(tax['on'])}  temas: {n_temas}  sub-temas: {n_subtemas}")
    for on_id, node in sorted(tax["on"].items()):
        print(f"  {on_id}: {node['label'][:70]}")
        for tema, tnode in node["temas"].items():
            print(f"    - {tema} ({len(tnode['subtemas'])} sub-temas)")
    print("\nCategorias excluidas (no son ON):",
          sorted({r['on'] for r in tax['excluded_categories']}))
    print("secundaria (Lineas Ceplan): ON=", len(tax["secondary"]))
